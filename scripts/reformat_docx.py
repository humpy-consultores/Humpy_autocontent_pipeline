"""
Reformatea Bioquimica_Tema_1_Enzimas.docx al formato estándar del pipeline:
- Heading 1  → MICROLECCIÓN N: Título
- Heading 2  → bloques temáticos principales
- List Bullet → ítems de lista cortos
- Normal      → cuerpo de texto explicativo
- Elimina quizzes (los genera el pipeline)
- Elimina emojis
"""

import re
from docx import Document

INPUT  = "Bioquimica_Tema_1_Enzimas.docx"
OUTPUT = "Bioquimica_Tema_1_Enzimas_FORMATO.docx"

# ---------------------------------------------------------------------------
# Emoji removal
# ---------------------------------------------------------------------------
EMOJI_RE = re.compile(
    "["
    "\U0001F300-\U0001F9FF"
    "\U0001FA00-\U0001FAFF"
    "☀-➿"
    "⌀-⏿"
    "⬀-⯿"
    "✂-➰"
    "︀-️"   # variation selectors
    "⃐-⃿"   # combining enclosing
    "️"
    "]+",
    flags=re.UNICODE,
)
NUMBER_EMOJI_RE = re.compile(r"^[0-9]+\s*[️⃣]?\s*")

MICROLECCION_RE = re.compile(r"microlección\s+(\d+)\s*[:\-]\s*(.+)", re.IGNORECASE)
QUIZ_RE         = re.compile(r"quiz interactivo", re.IGNORECASE)

# These patterns mark real section headers (Heading 2)
SECTION_HEADER_PATTERNS = [
    re.compile(r"^¿.+\?$"),                          # questions
    re.compile(r"^(tipos de|partes esenciales|casos específicos|forma de unión|"
               r"factores que afectan|hipótesis de|isoenzimas|sitio activo|"
               r"especificidad enzimática|inhibición reversible|inhibición irreversible|"
               r"enzimas alostéricas|efecto de la concentración|"
               r"características clave|ejemplo destacado|desnaturalización)", re.IGNORECASE),
    re.compile(r"^.{5,60}:$"),                        # short line ending in colon
]


def clean_emoji(text: str) -> str:
    text = EMOJI_RE.sub("", text)
    text = NUMBER_EMOJI_RE.sub("", text)
    return text.strip(" .:–-​")


def is_section_header(text: str) -> bool:
    for pat in SECTION_HEADER_PATTERNS:
        if pat.match(text):
            return True
    return False


def is_list_item(text: str) -> bool:
    # Short line, no sub-paragraphs, no long explanation
    return len(text) <= 100 and "\n" not in text and not text.endswith(".")


# ---------------------------------------------------------------------------
# Main reformatter
# ---------------------------------------------------------------------------

def reformat():
    src = Document(INPUT)
    dst = Document()

    dst.add_heading("BIOQUÍMICA – ENZIMAS", level=0)

    in_quiz = False

    for para in src.paragraphs:
        raw = para.text.strip()
        if not raw:
            continue

        # Detect quiz block start → skip until next microlección
        if QUIZ_RE.search(raw):
            in_quiz = True
            continue

        text = clean_emoji(raw)
        if not text:
            continue

        # Check microlección pattern before applying quiz skip
        ml_match = MICROLECCION_RE.match(text)
        if ml_match:
            in_quiz = False
            heading = f"MICROLECCIÓN {ml_match.group(1)}: {ml_match.group(2).strip()}"
            dst.add_heading(heading, level=1)
            continue

        if in_quiz:
            continue

        # Classify
        if is_section_header(text):
            dst.add_heading(text.rstrip(":"), level=2)
        elif is_list_item(text):
            dst.add_paragraph(text, style="List Bullet")
        else:
            dst.add_paragraph(text, style="Normal")

    dst.save(OUTPUT)
    print(f"Guardado: {OUTPUT}")


if __name__ == "__main__":
    reformat()
